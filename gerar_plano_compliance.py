from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document()

def h1(text):
    doc.add_heading(text, level=1)

def h2(text):
    doc.add_heading(text, level=2)

def p(text, bold=False, italic=False):
    par = doc.add_paragraph()
    run = par.add_run(text)
    run.bold = bold
    run.italic = italic
    return par

def bullet(text):
    doc.add_paragraph(text, style='List Bullet')

def table(headers, rows):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Light Grid Accent 1'
    hdr = t.rows[0].cells
    for i, htext in enumerate(headers):
        hdr[i].text = htext
        for run in hdr[i].paragraphs[0].runs:
            run.bold = True
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = val
    doc.add_paragraph("")

title = doc.add_heading("Plano de Implementação — Compliance SGSM", level=0)
sub = doc.add_paragraph("Auditoria de acesso → Criptografia em repouso → LGPD formal")
sub.runs[0].italic = True
meta = doc.add_paragraph("Uso interno  |  Baseado em Pendencias_Compliance_Escala_SGSM.docx (2026-08-21)  |  2026-08-25")
meta.runs[0].font.size = Pt(9)
meta.runs[0].font.color.rgb = RGBColor(0x66, 0x66, 0x66)

h1("Contexto")
p("O documento \"Pendencias_Compliance_Escala_SGSM.docx\" identificou 3 pendências antes do SGSM operar com dados reais de pacientes em produção, priorizadas nesta ordem: (1) auditoria de acesso a dados de paciente, (2) criptografia de CPF/dados sensíveis em repouso, (3) LGPD formal (consentimento, exportação, exclusão/anonimização). Este documento detalha como implementar cada item, com base em investigação direta do código-fonte atual dos repositórios sgsm, ms-sboot-auth, sgsm-ia e ms-whatsapp-bot.")

h2("Achados novos (não previstos no levantamento original)")
bullet("Bloqueador operacional: sgsm e ms-sboot-auth não têm Flyway (pom.xml sem a dependência, ddl-auto: none em todos os ambientes) — as migrations em db/migration/V003-V006 não rodam automaticamente. Precisa confirmar com o time como elas chegam ao banco de homologação/produção antes de escrever qualquer migration nova.")
bullet("Segunda cópia de PII não mencionada no levantamento original: MilvusIndexService.upsert() (sgsm-ia) grava o texto vetorizado completo (nome/CPF/e-mail em texto puro) também em crm.documento.conteudo no Postgres, não só no Milvus. Qualquer anonimização de paciente precisa purgar as duas cópias.")
bullet("VetorizacaoConsumer.processar() já lê o campo tipo/id da mensagem Redis mas nunca usa um campo operacao — hoje despacha para construir()+upsert() incondicionalmente. Um evento novo de anonimização exige um if/switch novo ali, não é só publicar o evento.")
bullet("Paciente.cpf é @Column(unique = true), texto puro, sem nenhuma criptografia.")
bullet("ContextoSeguranca expõe getPerfil()/getReferenciaId()/isPaciente()/etc., mas não getEmail(), embora o JwtAuthFilter já injete esse dado na requisição.")

h1("Ordem geral de implementação")
p("Auditoria → Criptografia → LGPD. Essa ordem é confirmada como correta porque criptografia e LGPD (exportação/anonimização) precisam registrar rastro de auditoria desde o primeiro dia em que essas rotas novas existirem; implementar auditoria por último seria retrabalho.")

table(
    ["Feature", "Ordem entre repositórios", "Motivo"],
    [
        ["Auditoria", "sgsm → ms-sboot-auth", "sgsm concentra o dado sensível (Paciente); ms-sboot-auth é greenfield de log (hoje zero logging estruturado)"],
        ["Criptografia", "só sgsm", "CPF só existe na entidade Paciente do sgsm; ms-sboot-auth não guarda CPF"],
        ["LGPD", "sgsm → sgsm-ia → ms-whatsapp-bot", "Endpoints de exportação/anonimização vivem no sgsm; purga do Milvus/crm.documento é o passo seguinte no sgsm-ia; consentimento no cadastro via bot é aditivo, pode vir em paralelo"],
    ]
)

h1("1. Auditoria de acesso a dados de paciente")
p("Decisão: tabela de audit log manual, chamada explícita nos services.", bold=True)
p("Não Hibernate Envers (audita mudança de estado, não leitura — que é o requisito mais cobrado em fiscalização de saúde). Não um novo Aspect/AOP (seria a primeira introdução de AOP no projeto, sem exemplo prévio; reavaliar depois que o padrão manual estiver rodando).")

h2("sgsm")
bullet("Nova tabela sgsm.log_acesso (migration nova, após resolver o bloqueador de Flyway): id, usuario_id, perfil, email, entidade, entidade_id, acao (LEITURA/CRIACAO/ATUALIZACAO/INATIVACAO/EXPORTACAO/ANONIMIZACAO), criado_em.")
bullet("Fechar o gap: adicionar getEmail() em ContextoSeguranca.java (espelhando getPerfil()).")
bullet("Novo AuditoriaService (registrar(entidade, entidadeId, acao)), assíncrono (@Async ou Redis Stream, no mesmo estilo do VetorizacaoPublisher já existente) para não acoplar latência/falha do log à transação de negócio.")
bullet("Instrumentar no MVP só PacienteService: consultar() (leitura — prioridade nº1), atualizar(), remover(), reativar(), e os métodos novos de LGPD (exportar(), anonimizar()).")

h2("ms-sboot-auth (greenfield — hoje zero log estruturado no serviço inteiro)")
bullet("Passo tático de baixo risco primeiro: adicionar Logger/log.info/log.warn em AuthService.login()/logout()/refresh()/registrar() (login sucesso/falha, refresh, detecção de reuso de refresh token revogado — hoje nada disso é logado).")
bullet("Evoluir para tabela auth.log_autenticacao depois, se o compliance exigir consulta estruturada em vez de grep de log.")

h1("2. Criptografia de CPF em repouso")
p("Decisão: AttributeConverter JPA (aplicação), não pgcrypto.", bold=True)
p("Chave gerenciada fora do SQL, mais fácil de rotacionar, consistente com o estilo do projeto. pgcrypto exigiria habilitar extensão nova e gerir chave dentro de query SQL.")

p("Conflito central a resolver:", bold=True)
p("cpf tem unique = true e é usado em checagem de duplicidade no cadastro (item já testado no QA funcional anterior como WB031). Criptografia não-determinística quebra as duas coisas.")

p("Recomendado: índice cego (blind index).", bold=True)
p("Manter cpf cifrado não-deterministicamente (mais seguro) numa coluna, adicionar cpf_hash (HMAC-SHA256 determinístico, só para busca/unicidade) com a constraint unique migrada para essa coluna. PacienteRepository.existsByCpf() passa a ser existsByCpfHash(hmac(cpf)).")
bullet("Alternativa mais rápida (MVP, menos segura): criptografia determinística direto na coluna cpf — mantém a constraint sem mudar schema, mas vaza informação de igualdade e é mais fraca contra ataque de dicionário (espaço de CPFs é finito/conhecido). Só usar se o prazo de compliance for curto, com plano explícito de migrar para blind index depois.")
bullet("Migração do dado existente: job único (não recorrente) que lê os CPFs em texto puro, calcula hash/cifra, popula as novas colunas, só então troca a constraint — roda antes do deploy do código novo, coordenado com o bloqueador de migrations.")
bullet("sgsm-ia: DocumentoBuilder.construirPaciente() lê cpf via JDBC direto e inclui no texto vetorizado — remover CPF do texto (não há motivo de negócio pro RAG \"saber\" o CPF) como parte deste item, antes mesmo do item LGPD.")

h1("3. LGPD formal")

h2("3.1 Consentimento no cadastro")
bullet("Paciente: campo consentimento_lgpd_em (timestamp).")
bullet("ms-whatsapp-bot (PacienteCreateRequest + executarCadastro()): novo campo consentimentoLgpd, nova pergunta explícita no fluxo conversacional antes de montar o payload — distinta de SessaoBot.consentimentoAceito (que é sobre uso do canal WhatsApp, não do cadastro clínico). É a mudança que mais toca o fluxo já testado em WB030 (ver seção de re-teste).")

h2("3.2 Exportação (portabilidade)")
bullet("Novo GET /v1/api/pacientes/{id}/exportar em PacienteController, reaproveitando o padrão self-only já usado em PacienteService.consultar() (paciente só exporta o próprio dado).")
bullet("Restringir a paciente-titular ou perfil administrativo (não médico/funcionário — exportação é direito do titular, não ferramenta operacional).")
bullet("Retorno: JSON com Paciente + Agendamentos do paciente. Registra auditoriaService.registrar(\"PACIENTE\", id, \"EXPORTACAO\").")

h2("3.3 Exclusão/anonimização")
bullet("Novo PATCH /v1/api/pacientes/{id}/anonimizar, mesmo padrão de reativar().")
bullet("PacienteService.anonimizar(): zera nome/cpf/email/telefone/endereço, mantém só o necessário para fins estatísticos/legais (conforme POLITICA_RETENCAO_DADOS.md), seta ativo=false.")
bullet("Novo campo anonimizado (boolean) em Paciente para distinguir \"inativo recuperável\" de \"anonimizado irreversível\" — bloquear reativar() quando anonimizado=true.")
bullet("Publica vetorizacaoPublisher.publicar(\"PACIENTE\", id, \"ANONIMIZAR\") (operação nova).")
bullet("VetorizacaoConsumer.processar() precisa passar a ler o campo operacao (hoje ignorado) e, quando for \"ANONIMIZAR\", chamar um novo método público em MilvusIndexService em vez de construir()+upsert() — e também purgar/anonimizar a linha correspondente em crm.documento.conteudo.")

h2("3.4 Job de expurgo automático")
bullet("Seguir o modelo já existente CrmAnaliticoScheduler.java (sgsm-ia): @Component com @Scheduled, injeta um @Service. Precisa de @EnableScheduling (ausente hoje no sgsm — seria o primeiro).")
bullet("Depende de campo novo encerrado_em (não existe ainda, só ativo), previsto na política de retenção mas nunca implementado.")
bullet("Não anonimiza automaticamente — gera relatório e aguarda aprovação manual (cron mensal, conforme POLITICA_RETENCAO_DADOS.md). O \"executar\" é uma chamada manual ao endpoint de anonimização (3.3), em lote.")

h1("Arquivos críticos")
bullet("sgsm/.../domain/Paciente.java — campo cpf (converter), consentimento_lgpd_em, anonimizado")
bullet("sgsm/.../security/ContextoSeguranca.java — getEmail()")
bullet("sgsm/.../service/PacienteService.java — pontos de auditoria, exportar(), anonimizar()")
bullet("sgsm/.../controller/PacienteController.java — GET /exportar, PATCH /anonimizar")
bullet("sgsm/.../repository/PacienteRepository.java — existsByCpfHash")
bullet("Novo: br.com.sgsm.service.AuditoriaService / br.com.sgsm.domain.LogAcesso + migration")
bullet("Novo: br.com.sgsm.converter.CpfCryptoConverter + cálculo de cpf_hash")
bullet("ms-sboot-auth/.../service/AuthService.java — logs de login/logout/refresh")
bullet("sgsm-ia/.../consumer/VetorizacaoConsumer.java — despachar por operacao")
bullet("sgsm-ia/.../service/MilvusIndexService.java — método público de remoção + purga de crm.documento")
bullet("sgsm-ia/.../service/DocumentoBuilder.java — remover CPF do texto vetorizado")
bullet("ms-whatsapp-bot/.../dto/PacienteCreateRequest.java + BotOrquestradorService.executarCadastro() — consentimento")

h1("Antes de codar: bloqueador a resolver")
p("Confirmar com o time/infra como as migrations V003-V006 chegam hoje ao banco de homologação/produção (não é via Spring Boot — sem Flyway, ddl-auto: none). Qualquer migration desta iniciativa (log_acesso, cpf_hash, consentimento_lgpd_em, anonimizado, encerrado_em) entra nesse mesmo processo — sem resolver isso primeiro, as migrations novas não têm caminho claro até o banco.")

h1("Impacto em testes")
table(
    ["Item existente", "Feature que impacta", "Risco", "Ação"],
    [
        ["WB030 (cadastro via bot)", "Consentimento (3.1)", "Alto — fluxo ganha pergunta obrigatória nova", "Re-executar do zero com o novo passo"],
        ["WB031 (duplicidade CPF)", "Criptografia (2)", "Alto — é exatamente o cenário que criptografia não-determinística sem blind index quebraria silenciosamente", "Re-executar após qualquer mudança na coluna cpf, incluindo variação de formatação"],
        ["WB041 (agendamento)", "Auditoria (1), indireto", "Baixo-médio — só quebra se a chamada de auditoria for síncrona e falhar revertendo a transação", "Validar que auditoria é best-effort/assíncrona; re-teste preventivo"],
    ]
)

p("Itens de teste novos (nunca cobertos, precisam de test-plan próprio):", bold=True)
bullet("Exportação de dados (caso feliz + acesso negado + completude do JSON)")
bullet("Anonimização (dado zerado no Postgres, evento publicado, vetor removido do Milvus, crm.documento purgado, reativar() bloqueado)")
bullet("Trilha de auditoria (log gerado a cada CRUD com usuário certo)")
bullet("Consentimento obrigatório no cadastro direto via sgsm (não só via bot)")
bullet("Job de expurgo (não anonimiza sem aprovação — comportamento esperado, não bug)")
p("Esta é a oportunidade de rodar a skill de QA funcional de backend pela primeira vez diretamente no sgsm (até agora só o ms-whatsapp-bot passou por ela) — as 3 features vivem majoritariamente no sgsm.")

h1("Verificação")
bullet("Compilar cada repositório após as mudanças (mvn compile) — ms-whatsapp-bot não tem suite de testes unitários, validação real é funcional.")
bullet("sgsm-ia e (se ganhar testes) sgsm/ms-sboot-auth: mvn test para regressão.")
bullet("Validação funcional real via skill de QA funcional de backend — reexecutar WB030/WB031/WB041 no ms-whatsapp-bot, e criar test-plan novo para o sgsm cobrindo auditoria/criptografia/LGPD.")

doc.save(r"C:\AmbienteDev\sgsm-front-medico\Plano_Implementacao_Compliance_SGSM.docx")
print("saved")
